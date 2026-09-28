# /// script
# requires-python = ">=3.12"
# dependencies = ["openpyxl", "psycopg[binary]", "httpx"]
# ///
"""Fills spirit.role and spirit.user_role from the owner's access sheet, and makes a Neon Auth user
for each person who has none.

    just spirit seed-access            # dry run against Neon: reads everything, writes nothing
    just spirit seed-access --apply    # creates the missing users, then writes both tables
    just spirit seed-access-local      # the throwaway database: roles, plus one per made-up user

The sheet stays out of git: it holds staff names and emails.
"""

import argparse
import secrets
import sys
import time
from dataclasses import dataclass, field
from pathlib import Path
from urllib.parse import urlsplit

import httpx
import openpyxl
import psycopg

REPO = Path(__file__).resolve().parents[4]
DEFAULT_SHEET = Path("/mnt/HDD/Downloads/Agent Access - review.xlsx")

# The sheet writes a group the way a person reads it; the database check constraint takes the
# names AccessGroups.cs uses.
GROUPS = {
    "customer/guest": "Guest",
    "guest": "Guest",
    "dealer": "Dealer",
    "tech service": "TechService",
    "inside sales": "InsideSales",
    "inside sales supervisor": "InsideSalesSupervisor",
    "tech service mgr": "TechServiceManager",
    "tech service manager": "TechServiceManager",
    "inside sales manager": "InsideSalesManager",
    "admin": "Admin",
}

# local-seed.sql's made-up users, keyed by the group each one stands for. The sheet gives some
# groups no role at all, so each gets a local-only role of its own.
LOCAL_USERS = {
    "Guest": "00000000-0000-4000-8000-000000000001",
    "Dealer": "00000000-0000-4000-8000-000000000002",
    "TechService": "00000000-0000-4000-8000-000000000003",
    "InsideSales": "00000000-0000-4000-8000-000000000004",
    "InsideSalesSupervisor": "00000000-0000-4000-8000-000000000005",
    "TechServiceManager": "00000000-0000-4000-8000-000000000006",
    "InsideSalesManager": "00000000-0000-4000-8000-000000000007",
    "Admin": "00000000-0000-4000-8000-000000000008",
}


@dataclass
class Person:
    name: str
    email: str
    here: bool
    roles: list[str]


@dataclass
class Report:
    created: list[str] = field(default_factory=list)
    found: list[str] = field(default_factory=list)
    to_create: list[str] = field(default_factory=list)
    stripped: list[str] = field(default_factory=list)
    no_group: list[str] = field(default_factory=list)

    def print(self, roles: dict[str, str | None], apply: bool) -> None:
        print(f"roles: {len(roles)}, with a group: {sum(1 for g in roles.values() if g)}")
        print("roles with no group: " + ", ".join(sorted(r for r, g in roles.items() if not g)))
        print(f"still here with no group, so no agent: {len(self.no_group)}")
        for email in self.no_group:
            print(f"  ? {email}")
        print(f"found on Neon: {len(self.found)}")
        if apply:
            print(f"created on Neon: {len(self.created)}")
            for email in self.created:
                print(f"  + {email}")
        else:
            print(f"would create on Neon: {len(self.to_create)}")
            for email in self.to_create:
                print(f"  + {email}")
        print(f"left the company but on Neon, roles {'removed' if apply else 'to remove'}: {len(self.stripped)}")
        for email in self.stripped:
            print(f"  - {email}")
        if not apply:
            print("\nDry run. Nothing was written. Run again with --apply.")


def read_sheet(path: Path) -> tuple[dict[str, str | None], list[Person]]:
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)

    roles: dict[str, str | None] = {}
    for name, label, *_ in book["Roles"].iter_rows(min_row=2, values_only=True):
        if not name:
            continue
        if label and label.strip().lower() not in GROUPS:
            sys.exit(f'The Roles tab gives {name} the group "{label}", which is none of: {", ".join(GROUPS)}.')
        roles[name.strip()] = GROUPS[label.strip().lower()] if label else None

    people = []
    for name, email, here, _check, _logins, listed, *_ in book["People"].iter_rows(min_row=2, values_only=True):
        if not email:
            continue
        person_roles = [r.strip() for r in (listed or "").split(",") if r.strip()]
        unknown = [r for r in person_roles if r not in roles]
        if unknown:
            sys.exit(f"{email} has roles the Roles tab does not list: {', '.join(unknown)}.")
        people.append(Person(name.strip(), email.strip().lower(), (here or "").strip().upper() == "Y", person_roles))

    return roles, people


def read_env(name: str) -> dict[str, str]:
    path = REPO / "secrets" / f"{name}.env"
    if not path.exists():
        sys.exit(f"{path} is missing. Run: just secrets init {name}")
    values = {}
    for line in path.read_text().splitlines():
        if "=" in line and not line.lstrip().startswith("#"):
            key, value = line.split("=", 1)
            values[key.strip()] = value.strip().strip('"')
    return values


def connect(npgsql: str) -> psycopg.Connection:
    """Opens the Npgsql-style connection string AgentCore reads, as libpq wants it."""
    parts = {k.strip().lower(): v.strip() for k, v in (p.split("=", 1) for p in npgsql.split(";") if "=" in p)}
    ssl = parts.get("ssl mode", "prefer").lower()
    return psycopg.connect(
        host=parts["host"],
        port=parts.get("port", "5432"),
        dbname=parts["database"],
        user=parts["username"],
        password=parts["password"],
        # libpq's verify-full needs a root certificate file; Neon's is signed by a public CA.
        sslmode="verify-full" if ssl == "verifyfull" else ssl.replace("verifyca", "verify-ca"),
        **({"sslrootcert": "system"} if ssl == "verifyfull" else {}),
    )


NO_TABLES = "spirit.user_role does not exist yet. Start the app once against this database so it migrates."


def has_tables(db: psycopg.Connection) -> bool:
    return db.execute("select to_regclass('spirit.user_role')").fetchone()[0] is not None


def check_tables(db: psycopg.Connection) -> None:
    if not has_tables(db):
        sys.exit(NO_TABLES)


class NeonAdmin:
    """Neon Auth's admin routes, behind the owner's session cookie (probe P3, 2026-09-27)."""

    def __init__(self, base_url: str, cookie: str) -> None:
        origin = "{0.scheme}://{0.netloc}".format(urlsplit(base_url))
        # create-user refuses a request with no Origin.
        self.http = httpx.Client(base_url=base_url.rstrip("/") + "/", headers={"Cookie": cookie, "Origin": origin}, timeout=30)

    def users(self) -> dict[str, str]:
        """Every user's id by email. One call per person trips Neon Auth's rate limit."""
        ids: dict[str, str] = {}
        while True:
            answer = self.send("GET", "admin/list-users", params={"limit": 100, "offset": len(ids)})
            self.raise_for(answer)
            page = answer.json()
            ids |= {u["email"].lower(): u["id"] for u in page["users"]}
            if not page["users"] or len(ids) >= page["total"]:
                return ids

    def create(self, person: Person) -> str:
        # Nobody keeps this password; people sign in with the email code.
        answer = self.send(
            "POST",
            "admin/create-user",
            json={"email": person.email, "name": person.name, "password": secrets.token_urlsafe(32)},
        )
        self.raise_for(answer)
        return answer.json()["user"]["id"]

    def send(self, method: str, path: str, **kwargs) -> httpx.Response:
        """Sends, and waits out Neon Auth's rate limit as long as it asks."""
        for _ in range(20):
            answer = self.http.request(method, path, **kwargs)
            if answer.status_code != 429:
                return answer
            wait = int(answer.headers.get("x-retry-after", "30")) + 1
            print(f"Neon Auth asks to slow down; waiting {wait}s.", file=sys.stderr)
            time.sleep(wait)
        return answer

    @staticmethod
    def raise_for(answer: httpx.Response) -> None:
        if answer.status_code == 401:
            sys.exit("Neon Auth refused the session. Copy a fresh NEON_ADMIN_COOKIE into secrets/dev.env.")
        if answer.status_code == 403:
            sys.exit(f"Neon Auth says the session is not an admin: {answer.text}")
        answer.raise_for_status()


def write_roles(db: psycopg.Connection, roles: dict[str, str | None]) -> None:
    db.cursor().executemany(
        "insert into spirit.role (name, access_group) values (%s, %s) "
        "on conflict (name) do update set access_group = excluded.access_group",
        list(roles.items()),
    )


def set_user_roles(db: psycopg.Connection, user_id: str, roles: list[str]) -> None:
    """Makes the person's roles exactly the given ones."""
    db.execute("delete from spirit.user_role where user_id = %s and not (role = any(%s))", (user_id, roles))
    db.cursor().executemany(
        "insert into spirit.user_role (user_id, role) values (%s, %s) on conflict do nothing",
        [(user_id, role) for role in roles],
    )


def seed_neon(args: argparse.Namespace, roles: dict[str, str | None], people: list[Person]) -> None:
    env = read_env(args.env)
    for key in ("POSTGRES_CONNECTION_STRING", "Auth__Neon__BaseUrl", "NEON_ADMIN_COOKIE"):
        if not env.get(key):
            sys.exit(f"secrets/{args.env}.env has no {key}.")

    admin = NeonAdmin(env["Auth__Neon__BaseUrl"], env["NEON_ADMIN_COOKIE"])
    report = Report()

    with connect(env["POSTGRES_CONNECTION_STRING"]) as db:
        if args.apply:
            check_tables(db)
        elif not has_tables(db):
            print(f"Warning: {NO_TABLES}\n")

        # Every read first, so a refused cookie stops the run before anything changes.
        known = admin.users()
        ids = {p.email: known.get(p.email) for p in people}

        for person in people:
            if person.here and not any(roles[r] for r in person.roles):
                report.no_group.append(person.email)
            if not person.here:
                if ids[person.email]:
                    report.stripped.append(person.email)
            elif ids[person.email]:
                report.found.append(person.email)
            else:
                report.to_create.append(person.email)

        if args.apply:
            for person in people:
                if person.here and not ids[person.email]:
                    ids[person.email] = admin.create(person)
                    report.created.append(person.email)

            with db.transaction():
                write_roles(db, roles)
                for person in people:
                    if ids[person.email]:
                        set_user_roles(db, ids[person.email], person.roles if person.here else [])

    report.print(roles, args.apply)


def seed_local(args: argparse.Namespace, roles: dict[str, str | None]) -> None:
    local_roles = {f"local-{group}": group for group in LOCAL_USERS}
    with connect(args.db) as db:
        check_tables(db)
        with db.transaction():
            write_roles(db, roles | local_roles)
            for group, user_id in LOCAL_USERS.items():
                set_user_roles(db, user_id, [f"local-{group}"])
    print(f"wrote {len(roles)} roles from the sheet and gave each of {len(LOCAL_USERS)} made-up users its local role.")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--sheet", type=Path, default=DEFAULT_SHEET)
    parser.add_argument("--apply", action="store_true", help="write; without it the run only reads")
    parser.add_argument("--env", default="dev", choices=["dev", "prod"], help="which secrets file names Neon")
    parser.add_argument("--local", metavar="CONNECTION", dest="db", help="seed this throwaway database instead of Neon")
    args = parser.parse_args()

    if not args.sheet.exists():
        sys.exit(f"{args.sheet} is missing. Pass --sheet.")

    roles, people = read_sheet(args.sheet)
    if args.db:
        seed_local(args, roles)
    else:
        seed_neon(args, roles, people)


if __name__ == "__main__":
    main()
